from pathlib import Path

from pptx import Presentation


def extract_slide_text(file_path: str | Path) -> list[dict[str, str | int]]:
    presentation = Presentation(str(file_path))
    slides: list[dict[str, str | int]] = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        text_parts = [
            shape.text.strip()
            for shape in slide.shapes
            if shape.has_text_frame and shape.text.strip()
        ]
        try:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        except (AttributeError, ValueError):
            notes = ""
        if notes:
            text_parts.append(notes)
        slides.append({"text": "\n".join(text_parts), "slide": slide_number})

    return slides